"""Pipeline local, explicable et non destructif d'analyse/traduction/reconstruction."""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from app.ai.ocr.paddle_local import PaddleOCRLocal
from app.ai.ocr.tesseract import TesseractOCR
from app.ai.quality.metrics import (
    protected_entity_issues,
    reference_is_valid,
    reference_metrics,
    token_delta,
)
from app.ai.translation.typography import typography_equivalent
from app.domain.classification import classify_document, detect_language

from .layout import detect_layout

SUPPORTED = {".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".docx", ".xlsx", ".pptx"}
OFFICE_FORMATS = {"xlsx", "pptx"}


def _warn(code, message, severity="warning", **extra):
    value = {"code": code, "severity": severity, "message": message}
    value.update(extra)
    return value


def _significant(text: str) -> bool:
    cleaned = re.sub(r"\s+", "", text or "")
    return len(cleaned) >= 4 and sum(c.isalnum() for c in cleaned) >= 3


def _requires_translation(text: str) -> bool:
    """Return false for standalone protected identifiers/numbers.

    Keeping an IBAN, date, amount, email or short proper-name invariant is not
    evidence that a translatable sentence was left untranslated.
    """
    value = (text or "").strip()
    if not value:
        return False
    if re.fullmatch(r"[\d\s.,:/+()\-]+", value, flags=re.UNICODE):
        return False
    if re.fullmatch(r"[\w.+-]+@[\w.-]+", value, flags=re.UNICODE):
        return False
    if re.fullmatch(r"(?=.*\d)[A-Z0-9._/+\-]+", value):
        return False
    if value.casefold() in {
        "iban",
        "bic",
        "swift",
        "eur",
        "usd",
        "mad",
        "signature",
        "email",
        "date",
        "total",
    }:
        return False
    return bool(re.search(r"[A-Za-zÀ-ÿ\u0600-\u06ff]", value))


def _intersects(a, b):
    if not a or not b or len(a) < 4 or len(b) < 4:
        return False
    return not (a[2] <= b[0] or a[0] >= b[2] or a[3] <= b[1] or a[1] >= b[3])


def _language_for(text):
    return detect_language(text)[0]


def _save_json(path: Path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _render_pix(page, path: Path, dpi=120):
    import fitz

    if page.rect.width * page.rect.height * (dpi / 72) ** 2 > 24_000_000:
        raise ValueError("PREPROCESSING_PIXEL_LIMIT")
    pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
    pix.save(str(path))
    return path


def _source_preview_image(image, path: Path):
    from PIL import Image

    Image.open(image).convert("RGB").save(path)


def _convert_docx_with_libreoffice(source: Path, output_dir: Path, timeout: float = 90.0) -> Path:
    """Convert DOCX with a disposable LibreOffice profile and a hard timeout."""
    soffice = "/usr/bin/soffice"
    if not Path(soffice).is_file():
        raise RuntimeError("LibreOffice absent: pagination DOCX non mesurable")
    output_dir.mkdir(parents=True, exist_ok=True)
    # The profile is isolated to prevent user extensions/macros and concurrent
    # workers from affecting the conversion. No document-provided macro is run.
    with (
        tempfile.TemporaryDirectory(prefix="lo-profile-") as profile,
        tempfile.TemporaryDirectory(prefix="lo-home-") as home,
    ):
        profile_url = Path(profile).as_uri()
        env = os.environ.copy()
        env.update({"HOME": home, "SAL_USE_VCLPLUGIN": "svp", "LANG": "C.UTF-8"})
        command = [
            soffice,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nofirststartwizard",
            "--norestore",
            "--nolockcheck",
            f"-env:UserInstallation={profile_url}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(source),
        ]
        try:
            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=max(1.0, min(float(timeout), 300.0)),
                check=False,
                env=env,
            )
        except subprocess.TimeoutExpired as exc:
            raise TimeoutError(f"LibreOffice timeout après {timeout:.1f}s") from exc
        if completed.returncode != 0:
            detail = (completed.stderr or completed.stdout or "échec sans sortie").strip()
            raise RuntimeError(f"LibreOffice conversion échouée: {detail[-500:]}")
    converted = output_dir / f"{source.stem}.pdf"
    if not converted.is_file() or converted.stat().st_size == 0:
        raise RuntimeError("LibreOffice n'a pas produit de PDF")
    return converted


def _docx_rendered_pages(
    source: Path, work: Path, options: dict
) -> tuple[list[dict], list[Path], list[dict]]:
    """Return real page dimensions/previews and conversion warnings."""
    warnings: list[dict] = []
    (work / "previews").mkdir(parents=True, exist_ok=True)
    try:
        rendered = _convert_docx_with_libreoffice(
            source, work / "libreoffice", options.get("libreoffice_timeout", 90)
        )
        import fitz

        doc = fitz.open(str(rendered))
        pages = []
        previews = []
        for number, page in enumerate(doc, 1):
            preview = work / "previews" / f"original-{number:03d}.png"
            _render_pix(page, preview)
            previews.append(preview)
            pages.append(
                {
                    "number": number,
                    "width": float(page.rect.width),
                    "height": float(page.rect.height),
                    "scanned": False,
                    "coordinates_kind": "pdf",
                    "blocks": [],
                }
            )
        doc.close()
        return pages, previews, warnings
    except TimeoutError as exc:
        warnings.append(_warn("DOCX_LIBREOFFICE_TIMEOUT", str(exc), "error"))
    except Exception as exc:
        warnings.append(_warn("DOCX_LIBREOFFICE_CONVERSION_FAILED", str(exc), "error"))
    return [], [], warnings


def _make_ocr(options: dict):
    backend = str(options.get("ocr_backend", "tesseract")).casefold()
    languages = options.get("ocr_languages") or options.get("languages")
    if backend in {"paddle", "paddleocr", "paddleocr-local"}:
        assets = options.get("ocr_assets") or options.get("paddle_assets")
        return PaddleOCRLocal(
            languages=languages, assets=assets, use_gpu=options.get("ocr_use_gpu", False)
        )
    if backend not in {"tesseract", "tesseract-local"}:
        raise ValueError(f"Backend OCR local inconnu: {backend}")
    return TesseractOCR(languages)


def _ocr_page(page, ocr, number):
    import fitz
    from PIL import Image

    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    result = ocr.image_to_blocks(image)
    sx, sy = page.rect.width / pix.width, page.rect.height / pix.height
    out = []
    for i, block in enumerate(result.blocks):
        x0, y0, x1, y1 = block["bbox"]
        text = block["text"]
        bbox = [
            max(0.0, min(float(page.rect.width), x0 * sx)),
            max(0.0, min(float(page.rect.height), y0 * sy)),
            max(0.0, min(float(page.rect.width), x1 * sx)),
            max(0.0, min(float(page.rect.height), y1 * sy)),
        ]
        out.append(
            {
                "id": f"p{number:03d}-b{i:04d}",
                "page": number,
                "kind": "text",
                "text": text,
                "bbox": bbox,
                "font_size": None,
                "confidence": block["confidence"],
                "language": _language_for(text),
                "provenance": "ocr:tesseract",
                "reading_order": i,
            }
        )
    out.sort(key=lambda block: (round(block["bbox"][1], 1), round(block["bbox"][0], 1)))
    for reading_order, block in enumerate(out):
        block["reading_order"] = reading_order
    return out, result.warnings


def _native_blocks(page, number):
    blocks = []
    data = page.get_text("dict")
    order = 0
    for raw in data.get("blocks", []):
        if raw.get("type") != 0:
            continue
        for line in raw.get("lines", []):
            text = "".join(span.get("text", "") for span in line.get("spans", []))
            bbox = line.get("bbox", raw.get("bbox"))
            if not text.strip():
                continue
            spans = [
                span
                for line2 in raw.get("lines", [])
                for span in line2.get("spans", [])
                if span.get("text", "").strip()
            ]
            font_size = (
                sum(float(span.get("size", 0)) for span in spans) / len(spans) if spans else None
            )
            blocks.append(
                {
                    "id": f"p{number:03d}-b{order:04d}",
                    "page": number,
                    "kind": "text",
                    "text": text.strip(),
                    "bbox": [float(x) for x in bbox],
                    "font_size": font_size,
                    "confidence": 1.0,
                    "language": _language_for(text),
                    "provenance": "native",
                    "reading_order": order,
                }
            )
            order += 1
    blocks.sort(
        key=lambda block: (round(float(block["bbox"][1]), 1), round(float(block["bbox"][0]), 1))
    )
    for reading_order, block in enumerate(blocks):
        block["reading_order"] = reading_order
    return blocks


def _tables(page, number, blocks=None, image=None):
    """Return structural candidates and layout coverage without duplicating text blocks."""
    sx = float(page.rect.width) / image.width if image is not None else 1.0
    sy = float(page.rect.height) / image.height if image is not None else 1.0
    layout_blocks = [
        {
            **block,
            "bbox": [
                block["bbox"][0] / sx,
                block["bbox"][1] / sy,
                block["bbox"][2] / sx,
                block["bbox"][3] / sy,
            ],
        }
        if image is not None and block.get("bbox")
        else block
        for block in (blocks or [])
    ]
    result = detect_layout(
        image,
        page=page if image is None else None,
        blocks=layout_blocks,
        page_number=number,
        width=float(page.rect.width),
        height=float(page.rect.height),
    )
    if image is not None:

        def source_coordinates(value):
            if isinstance(value, dict):
                if isinstance(value.get("bbox"), (list, tuple)) and len(value["bbox"]) == 4:
                    box = value["bbox"]
                    value["bbox"] = [box[0] * sx, box[1] * sy, box[2] * sx, box[3] * sy]
                    value["coordinates_kind"] = "source"
                for nested in value.values():
                    source_coordinates(nested)
            elif isinstance(value, list):
                for nested in value:
                    source_coordinates(nested)

        source_coordinates(result)
    tables = result.get("tables", [])
    # Merge P1/P2: enrichment runs on SOURCE (PDF point) coordinates only.
    from .pdf_fidelity import enrich_tables

    enrich_tables(page, tables, blocks or [])
    return tables, result


def _clip_candidate_regions(regions, pages, warnings):
    """Clamp AUTOMATIC preprocessing candidates to page bounds (float rounding).

    User annotations are not passed here: they still fail closed in
    normalize_regions. Degenerate candidates are dropped with a warning.
    """
    sizes = {int(p["number"]): (float(p["width"]), float(p["height"])) for p in pages}
    kept = []
    for region in regions or []:
        page = int(region.get("page", 1))
        box = region.get("bbox") or []
        if page not in sizes or len(box) != 4:
            warnings.append(
                _warn(
                    "PROTECTED_CANDIDATE_INVALID",
                    "Candidat visuel automatique sans page/bbox valide ignoré; revue requise.",
                    "warning",
                    page=page,
                )
            )
            continue
        width, height = sizes[page]
        clipped = [
            min(max(float(box[0]), 0.0), width),
            min(max(float(box[1]), 0.0), height),
            min(max(float(box[2]), 0.0), width),
            min(max(float(box[3]), 0.0), height),
        ]
        if clipped[2] - clipped[0] < 0.5 or clipped[3] - clipped[1] < 0.5:
            warnings.append(
                _warn(
                    "PROTECTED_CANDIDATE_DEGENERATE",
                    "Candidat visuel automatique dégénéré ignoré; revue requise.",
                    "warning",
                    page=page,
                )
            )
            continue
        kept.append(
            {
                **region,
                "bbox": clipped,
                "provenance": region.get("provenance") or "automatic_candidate:preprocessing",
            }
        )
    return kept


def _iou(a, b):
    ix = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    union = (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - inter
    return inter / union if union > 0 else 0.0


def _exclude_table_art(regions, page_tables, warnings, page_number, threshold=0.6):
    """Drop automatic raster candidates that coincide with a detected table frame.

    Only geometry equal to a table (IoU >= threshold) is excluded; a stamp that
    partially overlaps a table keeps its protection. Every exclusion is traced.
    """
    boxes = [t.get("bbox") for t in page_tables or [] if t.get("bbox") and len(t["bbox"]) == 4]
    kept = []
    for region in regions or []:
        box = region.get("bbox") or []
        match = next((b for b in boxes if len(box) == 4 and _iou(box, b) >= threshold), None)
        if match is not None:
            warnings.append(
                _warn(
                    "PROTECTED_CANDIDATE_TABLE_ART",
                    "Candidat visuel automatique identique au cadre d'un tableau détecté: non protégé, tableau traité par le renderer de cellules.",
                    "info",
                    page=page_number,
                    bbox=[round(float(v), 2) for v in box],
                )
            )
            continue
        kept.append(region)
    return kept


def _apply_protection(pages, protected):
    for page in pages:
        for block in page["blocks"]:
            for zone in protected:
                if int(zone.get("page", 0)) == page["number"] and _intersects(
                    block["bbox"], zone.get("bbox", [])
                ):
                    block["protected"] = True


def _docx_text(element):
    from docx.oxml.ns import qn

    return "".join(node.text or "" for node in element.iter() if node.tag == qn("w:t"))


def _docx_paragraph_block(text, block_id, order, source_ref, kind="paragraph"):
    return {
        "id": block_id,
        "page": 1,
        "kind": kind,
        "text": text,
        "bbox": [36, 36 + order * 16, 576, 50 + order * 16],
        "coordinates_kind": "logical",
        "artificial_bbox": True,
        "font_size": None,
        "confidence": 1.0,
        "language": _language_for(text),
        "provenance": "docx:xml",
        "reading_order": order,
        "source_ref": source_ref,
    }


def _docx_ir(path: Path, work: Path, options):
    from .docx_analysis import collect_docx_content

    rendered_pages, _previews, conversion_warnings = _docx_rendered_pages(path, work, options)
    page = (
        rendered_pages[0]
        if rendered_pages
        else {
            "number": 1,
            "width": 612.0,
            "height": 792.0,
            "scanned": False,
            "coordinates_kind": "logical-fallback",
            "blocks": [],
        }
    )
    blocks, tables_ir, hierarchy = collect_docx_content(path, _docx_paragraph_block)
    header_footer_refs = [b for b in blocks if b["source_ref"]["xml_part"] != "word/document.xml"]

    source_copy = work / "source" / path.name
    source_copy.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, source_copy)
    xml_hash = hashlib.sha256(source_copy.read_bytes()).hexdigest()
    text = "\n".join(block["text"] for block in blocks)
    lang, languages, confidence = detect_language(text)
    cls = classify_document(text, languages)
    warnings = list(conversion_warnings)
    warnings.extend(cls["warnings"])
    # Map logical text to the actual rendered page where possible. Coordinates
    # remain evidence from LibreOffice; unmatched blocks are review-only.
    pages = rendered_pages or [page]
    rendered_pdf = next((work / "libreoffice").glob("*.pdf"), None)
    if rendered_pdf:
        import fitz

        rendered_doc = fitz.open(str(rendered_pdf))
        for block in blocks:
            match = None
            needle = re.sub(r"\s+", " ", block["text"]).strip().casefold()
            for page_number, rendered_page in enumerate(rendered_doc, 1):
                for raw in rendered_page.get_text("blocks"):
                    candidate = re.sub(r"\s+", " ", raw[4]).strip().casefold()
                    if needle and (needle in candidate or candidate in needle):
                        match = (page_number, list(raw[:4]))
                        break
                if match:
                    break
            if match:
                block["page"], block["bbox"] = match
                block["coordinates_kind"] = "pdf"
                block["artificial_bbox"] = False
            else:
                block["unmapped_in_render"] = True
        rendered_doc.close()
    for rendered_page in pages:
        rendered_page["blocks"] = [b for b in blocks if b.get("page") == rendered_page["number"]]
    if not rendered_pages:
        warnings.append(
            _warn(
                "DOCX_PAGINATION_UNAVAILABLE",
                "La pagination réelle est indisponible; aucune preview synthétique n'est présentée comme un rendu.",
                "error",
            )
        )
    if header_footer_refs:
        warnings.append(
            _warn(
                "DOCX_HEADER_FOOTER_PRESENT",
                "En-têtes/pieds détectés et adressables par source_ref; leur résultat doit être vérifié visuellement.",
                "warning",
            )
        )
    return {
        "schema_version": "1.0",
        "source_path": str(path.resolve()),
        "format": "docx",
        "language": lang,
        "languages": languages,
        "language_confidence": confidence,
        "document_type": cls["document_type"],
        "document_type_confidence": cls["confidence"],
        "domain": cls["domain"],
        "pages": pages,
        "tables": tables_ir,
        **hierarchy,
        "warnings": warnings,
        "versions": {
            "source_xml_sha256": xml_hash,
            "source_copy": str(source_copy),
            "docx_mapping": "package-w:t-proportional-v3",
            "pagination": "libreoffice-pdf" if rendered_pages else None,
        },
        "classifier": cls,
    }


def _office_ir(path: Path, work: Path, options):
    """RUN28 route A: native XLSX/PPTX IR (logical units = sheets / slides)."""
    from .ooxml_office import VERSION as OFFICE_VERSION
    from .ooxml_office import collect_office_content

    fmt = path.suffix.lower().lstrip(".")
    content = collect_office_content(path)
    source_copy = work / "source" / path.name
    source_copy.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, source_copy)
    blocks = []
    for order, item in enumerate(content["blocks"]):
        blocks.append(
            {
                "id": item["id"],
                "page": item["page"],
                "kind": item["kind"],
                "text": item["text"],
                "bbox": [36, 36 + (order % 40) * 16, 576, 50 + (order % 40) * 16],
                "coordinates_kind": "logical",
                "artificial_bbox": True,
                "font_size": None,
                "confidence": 1.0,
                "language": _language_for(item["text"]),
                "provenance": f"{fmt}:xml",
                "reading_order": order,
                "source_ref": item["source_ref"],
                **({"autofit": item["autofit"]} if "autofit" in item else {}),
            }
        )
    pages = [
        {
            "number": unit["number"],
            "width": 612.0,
            "height": 792.0,
            "scanned": False,
            "coordinates_kind": "logical-fallback",
            "unit": "sheet" if fmt == "xlsx" else "slide",
            "hidden": unit["hidden"],
            "blocks": [b for b in blocks if b["page"] == unit["number"]],
        }
        for unit in content["units"]
    ]
    text = "\n".join(block["text"] for block in blocks)
    lang, languages, confidence = detect_language(text)
    cls = classify_document(text, languages)
    warnings = list(cls["warnings"])
    inventory = content["inventory"]
    warnings.append(
        _warn(
            "OFFICE_VISUAL_PROOF_UNAVAILABLE",
            "Aucune preuve visuelle mesurée pour ce format : vérification native du paquet "
            "seulement (largeurs de cellules / ajustement du texte non mesurés).",
            "warning",
        )
    )
    untranslated = inventory.get("untranslated_text_parts") or {}
    if untranslated:
        warnings.append(
            _warn(
                "OFFICE_TEXT_PARTS_NOT_TRANSLATED",
                "Textes hors périmètre conservés tels quels (commentaires, notes, graphiques, SmartArt).",
                "warning",
                parts=len(untranslated),
                text_nodes=sum(max(0, n) for n in untranslated.values()),
            )
        )
    if inventory.get("header_footer_text"):
        warnings.append(
            _warn(
                "XLSX_HEADER_FOOTER_NOT_TRANSLATED",
                "En-têtes/pieds d'impression conservés tels quels (codes de mise en page).",
                "warning",
            )
        )
    if inventory.get("formula_text_cells"):
        warnings.append(
            _warn(
                "XLSX_FORMULA_TEXT_NOT_TRANSLATED",
                "Résultats texte de formules conservés (recalculés par le tableur).",
                "info",
                cells=inventory["formula_text_cells"],
            )
        )
    if inventory.get("hidden_units"):
        warnings.append(
            _warn(
                "OFFICE_HIDDEN_UNITS_PRESENT",
                "Feuilles ou diapositives masquées présentes : traduites, à vérifier.",
                "warning",
                hidden=inventory["hidden_units"],
            )
        )
    previews = []
    if options.get("office_previews", True):
        rendered_pages, previews, conversion_warnings = _docx_rendered_pages(path, work, options)
        for warning in conversion_warnings:
            warning["code"] = warning["code"].replace("DOCX_", "OFFICE_")
            warning["severity"] = "warning"
        warnings.extend(conversion_warnings)
    return {
        "schema_version": "1.0",
        "source_path": str(path.resolve()),
        "format": fmt,
        "language": lang,
        "languages": languages,
        "language_confidence": confidence,
        "document_type": cls["document_type"],
        "document_type_confidence": cls["confidence"],
        "domain": cls["domain"],
        "pages": pages,
        "tables": [],
        "warnings": warnings,
        "office_inventory": {k: v for k, v in inventory.items() if not k.endswith("_digest")},
        "versions": {
            "source_copy": str(source_copy),
            "office_mapping": OFFICE_VERSION,
            "preview": "libreoffice-pdf" if previews else None,
        },
        "classifier": cls,
    }


def _analyze_image(image, ocr, number, work, width, height, options, style_image=None):
    from .preprocessing import map_bbox, preprocess_image

    sx, sy = width / image.width, height / image.height
    supplied = []
    for region in options.get("protected_regions") or []:
        if int(region.get("page", 0)) == number:
            supplied.append(
                {
                    **region,
                    "bbox": [
                        region["bbox"][0] / sx,
                        region["bbox"][1] / sy,
                        region["bbox"][2] / sx,
                        region["bbox"][3] / sy,
                    ],
                }
            )
    processed = preprocess_image(
        image,
        ocr,
        work / "preprocessing" / f"page-{number:03d}",
        page=number,
        profile=options.get(
            "preprocessing_profile", options.get("image_preprocessing_profile", "auto")
        ),
        protected_regions=supplied,
        max_variants=options.get("preprocessing_max_variants", 6),
        orientation=options.get("preprocessing_orientation", True),
    )
    processed.image.save(work / "previews" / f"enhanced-{number:03d}.png")
    # Style evidence (stroke, ink and paper colour) is measured on the ORIGINAL
    # raster, never on the enhanced OCR copy.
    import numpy as np

    from .segmentation import line_style_evidence

    # The 2x analysis raster (144 dpi) saturates sub-pixel strokes of 7-10 pt
    # text; scans are 200-300 dpi, so style is measured on a 4x render.
    style_source = style_image if style_image is not None else image
    original_rgb = np.asarray(style_source.convert("RGB"))
    style_scale = style_source.width / float(width)
    blocks = []
    for i, block in enumerate(processed.ocr.blocks):
        bbox = map_bbox(
            block["bbox"],
            processed.metadata["transform_processing_to_source"],
            image.width,
            image.height,
        )
        item = {
            "id": f"p{number:03d}-b{i:04d}",
            "page": number,
            "kind": "text",
            "text": block["text"],
            "bbox": [bbox[0] * sx, bbox[1] * sy, bbox[2] * sx, bbox[3] * sy],
            "font_size": None,
            "confidence": block["confidence"],
            "language": _language_for(block["text"]),
            "provenance": "ocr:" + getattr(processed.ocr, "backend", "local"),
            "reading_order": i,
            "coordinates_kind": "source",
        }
        style_box = [
            bbox[0] * sx * style_scale,
            bbox[1] * sy * style_scale,
            bbox[2] * sx * style_scale,
            bbox[3] * sy * style_scale,
        ]
        evidence = line_style_evidence(original_rgb, style_box, points_per_pixel=1.0 / style_scale)
        if evidence is not None:
            item["style_evidence"] = evidence
        blocks.append(item)
    regions = [
        {
            **r,
            "bbox": [r["bbox"][0] * sx, r["bbox"][1] * sy, r["bbox"][2] * sx, r["bbox"][3] * sy],
            "coordinates_kind": "source",
        }
        for r in processed.protected_visual_regions
    ]
    # Renderer operates in PDF points or original image pixels. Composed matrices
    # explicitly convert those SOURCE coordinates to/from processing-image pixels.
    scale = np.diag([sx, sy, 1.0])
    meta = processed.metadata
    meta["processing_image_path"] = str(work / "previews" / f"enhanced-{number:03d}.png")
    meta["transform_processing_to_source_pixels"] = meta["transform_processing_to_source"]
    meta["transform_source_pixels_to_processing"] = meta["transform_source_to_processing"]
    meta["transform_processing_to_source"] = (
        scale @ np.asarray(meta["transform_processing_to_source"])
    ).tolist()
    meta["transform_source_to_processing"] = (
        np.asarray(meta["transform_source_to_processing"]) @ np.linalg.inv(scale)
    ).tolist()
    meta["source_coordinate_size"] = [width, height]
    return blocks, meta, regions, list(processed.ocr.warnings) + processed.warnings


def analyze_document(source_path: str, work_dir: str, options: dict | None = None) -> dict:
    options = options or {}
    path = Path(source_path).expanduser().resolve()
    work = Path(work_dir).resolve()
    work.mkdir(parents=True, exist_ok=True)
    if not path.is_file():
        raise FileNotFoundError(source_path)
    if path.suffix.lower() not in SUPPORTED:
        raise ValueError(f"Format non supporté: {path.suffix}")
    source_sha = hashlib.sha256(path.read_bytes()).hexdigest()
    (work / "previews").mkdir(exist_ok=True)
    if path.suffix.lower().lstrip(".") in OFFICE_FORMATS:
        ir = _office_ir(path, work, options)
        ir["versions"]["source_sha256"] = source_sha
        ir["source_metadata"] = {
            "name": path.name,
            "size": path.stat().st_size,
            "format": ir["format"],
            "version": 1,
        }
        if hashlib.sha256(path.read_bytes()).hexdigest() != source_sha:
            raise RuntimeError("ORIGINAL_INTEGRITY_MISMATCH")
        _save_json(work / "analysis.json", ir)
        return ir
    if path.suffix.lower() == ".docx":
        ir = _docx_ir(path, work, options)
        from .protected_visuals import protect_ir

        protect_ir(ir, options.get("protected_visual_regions") or options.get("protected_regions"))
        ir["versions"]["source_sha256"] = source_sha
        ir["source_metadata"] = {
            "name": path.name,
            "size": path.stat().st_size,
            "format": "docx",
            "version": 1,
        }
        if hashlib.sha256(path.read_bytes()).hexdigest() != source_sha:
            raise RuntimeError("ORIGINAL_INTEGRITY_MISMATCH")
        _save_json(work / "analysis.json", ir)
        return ir
    import fitz
    from PIL import Image, ImageSequence

    from .preprocessing import VERSION, protected_candidates

    ocr = _make_ocr(options)
    warnings, pages, tables, all_text, protected = [], [], [], [], []
    max_pages = int(options.get("max_pages", 200))
    metadata = {}
    if path.suffix.lower() == ".pdf":
        doc = fitz.open(str(path))
        try:
            if len(doc) > max_pages:
                raise ValueError("PAGE_LIMIT")
            metadata = doc.metadata
            for number, page_obj in enumerate(doc, 1):
                native = _native_blocks(page_obj, number)
                scanned = not _significant("\n".join(b["text"] for b in native))
                preprocessing = None
                regions = []
                page_blocks = native
                if page_obj.rect.width * page_obj.rect.height * (4 if scanned else 1) > 24_000_000:
                    raise ValueError("PREPROCESSING_PIXEL_LIMIT")
                pix = page_obj.get_pixmap(
                    matrix=fitz.Matrix(2 if scanned else 1, 2 if scanned else 1), alpha=False
                )
                image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                if scanned:
                    style_pix = page_obj.get_pixmap(matrix=fitz.Matrix(4, 4), alpha=False)
                    style_image = Image.frombytes(
                        "RGB", [style_pix.width, style_pix.height], style_pix.samples
                    )
                    page_blocks, preprocessing, regions, page_warnings = _analyze_image(
                        image,
                        ocr,
                        number,
                        work,
                        float(page_obj.rect.width),
                        float(page_obj.rect.height),
                        options,
                        style_image=style_image,
                    )
                    warnings.extend(page_warnings)
                    warnings.extend(
                        [
                            _warn(
                                "OCR_LIMITATION",
                                "OCR local: exactitude non validée sans transcription humaine.",
                                page=number,
                            ),
                            _warn(
                                "NON_TEXTUAL_UNIDENTIFIED",
                                "Régions visuelles candidates non exhaustives; revue requise.",
                                page=number,
                            ),
                            _warn(
                                "SIGNATURE_AUTHENTICITY_NOT_VERIFIED",
                                "Aucune authentification des signatures/cachets.",
                                page=number,
                            ),
                        ]
                    )
                if not scanned:
                    import numpy as np

                    regions = protected_candidates(np.asarray(image), number)
                    sx, sy = (
                        float(page_obj.rect.width) / image.width,
                        float(page_obj.rect.height) / image.height,
                    )
                    for region in regions:
                        b = region["bbox"]
                        region["bbox"] = [b[0] * sx, b[1] * sy, b[2] * sx, b[3] * sy]
                        region["coordinates_kind"] = "source"
                if not page_blocks:
                    warnings.append(
                        _warn(
                            "BLANK_OR_UNREADABLE_PAGE",
                            "Page conservée sans texte exploitable.",
                            page=number,
                        )
                    )
                try:
                    if list(page_obj.widgets() or []):
                        warnings.append(
                            _warn(
                                "PDF_SIGNATURE_WIDGET_REQUIRES_REVIEW",
                                "Widget PDF: aucune validation cryptographique du dérivé.",
                                page=number,
                            )
                        )
                except Exception:
                    pass
                # Layout stays on ORIGINAL raster and source-coordinate OCR boxes.
                page_tables, layout = _tables(
                    page_obj, number, page_blocks, image if scanned else None
                )
                if not scanned:
                    # Native PDF: vector table art is handled by the P2 renderer;
                    # a raster candidate equal to a detected table is not a stamp.
                    regions = _exclude_table_art(regions, page_tables, warnings, number)
                pages.append(
                    {
                        "number": number,
                        "width": float(page_obj.rect.width),
                        "height": float(page_obj.rect.height),
                        "scanned": scanned,
                        "blocks": page_blocks,
                        "layout": layout,
                        "preprocessing": preprocessing,
                        "processing_image_path": preprocessing.get("processing_image_path")
                        if preprocessing
                        else None,
                        "protected_visual_regions": regions,
                    }
                )
                protected.extend(regions)
                tables.extend(page_tables)
                all_text.extend(b["text"] for b in page_blocks)
                _render_pix(page_obj, work / "previews" / f"original-{number:03d}.png")
        finally:
            doc.close()
        fmt = "pdf"
    else:
        with Image.open(path) as source_image:
            if getattr(source_image, "n_frames", 1) > max_pages:
                raise ValueError("PAGE_LIMIT")
            for number, frame in enumerate(ImageSequence.Iterator(source_image), 1):
                if frame.width * frame.height > 24_000_000:
                    raise ValueError("PREPROCESSING_PIXEL_LIMIT")
                image = frame.convert("RGB")
                width, height = image.size
                page_blocks, preprocessing, regions, page_warnings = _analyze_image(
                    image, ocr, number, work, width, height, options
                )
                warnings.extend(page_warnings)
                layout = detect_layout(
                    image,
                    blocks=page_blocks,
                    page_number=number,
                    width=float(width),
                    height=float(height),
                )
                tables.extend(layout.get("tables", []))
                pages.append(
                    {
                        "number": number,
                        "width": float(width),
                        "height": float(height),
                        "scanned": True,
                        "blocks": page_blocks,
                        "layout": layout,
                        "preprocessing": preprocessing,
                        "processing_image_path": preprocessing.get("processing_image_path")
                        if preprocessing
                        else None,
                        "protected_visual_regions": regions,
                    }
                )
                protected.extend(regions)
                all_text.extend(b["text"] for b in page_blocks)
                image.save(work / "previews" / f"original-{number:03d}.png")
        warnings.append(
            _warn("OCR_LIMITATION", "OCR local: exactitude non validée sans transcription humaine.")
        )
        fmt = path.suffix.lower().lstrip(".")
    _apply_protection(pages, [*(options.get("protected_regions") or []), *protected])
    text = "\n".join(all_text)
    lang, languages, confidence = detect_language(text)
    cls = classify_document(text, languages)
    warnings.extend(cls["warnings"])
    if len(languages) > 1:
        warnings.append(
            _warn("MIXED_LANGUAGES", "Distribution heuristique des langues à vérifier.")
        )
    ir = {
        "schema_version": "1.0",
        "source_path": str(path),
        "format": fmt,
        "language": lang,
        "languages": languages,
        "language_confidence": confidence,
        "document_type": cls["document_type"],
        "document_type_confidence": cls["confidence"],
        "domain": cls["domain"],
        "pages": pages,
        "tables": tables,
        "protected_visual_regions": _clip_candidate_regions(protected, pages, warnings),
        "warnings": warnings,
        "versions": {
            "analysis": "local-ocr-native-v3",
            "source_sha256": source_sha,
            "preprocessing": VERSION,
        },
        "source_metadata": {
            "name": path.name,
            "format": fmt,
            "size": path.stat().st_size,
            "pages": len(pages),
            "metadata": metadata,
            "version": 1,
        },
        "timings": {
            "ocr_seconds": sum(
                (p.get("preprocessing") or {}).get("duration_seconds", 0) for p in pages
            )
        },
        "classifier": cls,
    }
    # P2 merges P1 automatic candidates, PDF source visuals and user annotations.
    from .protected_visuals import protect_ir

    protect_ir(ir, options.get("protected_visual_regions") or options.get("protected_regions"))
    if options.get("ocr_segmentation", True):
        # Physical OCR lines -> logical translation units (paragraphs, list
        # items, multi-line titles, cell contents), with line-level traceability.
        from .segmentation import segment_ir

        segment_ir(ir, detect_language=_language_for)
    if options.get("line_verification", True) and any(p.get("scanned") for p in ir["pages"]):
        run_line_verification(ir, path)
    if (
        hashlib.sha256(path.read_bytes()).hexdigest() != source_sha
        or ir["versions"]["source_sha256"] != source_sha
    ):
        raise RuntimeError("ORIGINAL_INTEGRITY_MISMATCH")
    _save_json(work / "analysis.json", ir)
    return ir


def run_line_verification(ir, source):
    """Highlight evidence + critical-entity double reading on the original raster.

    A failure never passes silently: the IR records NOT_EXECUTED and a review
    warning, so the missing safety net is visible downstream."""
    from .line_verification import verify_lines

    try:
        return verify_lines(ir, source)
    except Exception as exc:  # noqa: BLE001 - recorded, never hidden
        ir["line_verification"] = {"status": "NOT_EXECUTED", "error": type(exc).__name__}
        ir.setdefault("warnings", []).append(
            {
                "code": "LINE_VERIFICATION_NOT_EXECUTED",
                "severity": "warning",
                "message": "Double lecture des entités et mesure du surlignage non exécutées: revue humaine requise.",
            }
        )
        return ir["line_verification"]


def _all_blocks(ir):
    return [block for page in ir.get("pages", []) for block in page.get("blocks", [])]


def _source_text(ir):
    return "\n".join(block.get("text", "") for block in _all_blocks(ir))


def _render_diff(source, candidate, out):
    from PIL import Image, ImageChops, ImageEnhance

    first = Image.open(source).convert("RGB")
    second = Image.open(candidate).convert("RGB")
    if first.size != second.size:
        second = second.resize(first.size)
    diff = ImageEnhance.Contrast(ImageChops.difference(first, second)).enhance(4)
    diff.save(out)
    # This is a raw normalized pixel difference, not a visual similarity score.
    return sum(diff.convert("L").getdata()) / (255.0 * diff.width * diff.height)


def _preview_pdf(path, work, prefix):
    import fitz

    (Path(work) / "previews").mkdir(parents=True, exist_ok=True)
    doc = fitz.open(str(path))
    outputs = []
    for number, page in enumerate(doc, 1):
        output = Path(work) / "previews" / f"{prefix}-{number:03d}.png"
        _render_pix(page, output)
        outputs.append(output)
    doc.close()
    return outputs


def _protected_for(block, protected_regions):
    # Run18 (F15): an automatic check-box glyph never protects its block (the
    # renderer carves it out of masks). Counting it here classified a block
    # returned unchanged by the model as "protected, kept" (warning) instead
    # of UNTRANSLATED (error), hiding an untranslated label.
    from .protected_visuals import carves_only

    return block.get("protected") or any(
        int(zone.get("page", 0)) == block.get("page")
        and not carves_only(zone)
        and _intersects(block.get("bbox", []), zone.get("bbox", []))
        for zone in protected_regions
    )


def _docx_text_nodes(element):
    from docx.oxml.ns import qn

    return [node for node in element.iter() if node.tag == qn("w:t")]


def _replace_docx_text(element, target):
    from .docx_fidelity import redistribute_text

    result = redistribute_text(element, target)
    return result["node_count"] if result["applied"] else False


def _docx_paragraph_for_ref(doc, ref):
    """Resolve a paragraph without flattening the surrounding DOCX structure."""
    if ref.get("header_footer"):
        section = doc.sections[int(ref["section_index"])]
        collection = section.header if ref["header_footer"] == "header" else section.footer
        return collection.paragraphs[int(ref["paragraph_index"])]
    if "paragraph_index" in ref and "table_index" not in ref:
        return doc.paragraphs[int(ref["paragraph_index"])]
    cell = (
        doc.tables[int(ref["table_index"])]
        .rows[int(ref["row_index"])]
        .cells[int(ref["cell_index"])]
    )
    return cell.paragraphs[int(ref.get("cell_paragraph_index", 0))]


def _with_render_proof(ir, translations, output, result, protected_regions):
    from .render_metrics import measure_render_fidelity

    render_source = ir["source_path"]
    if ir.get("format") in {"tif", "tiff"}:
        from .raster_adapter import raster_pdf_copy

        render_source = str(
            raster_pdf_copy(render_source, output.parent / "source-raster-copy.pdf")
        )
    pages = {page["number"]: page for page in ir.get("pages", [])}
    regions = [
        {
            **region,
            "page_width": pages.get(region.get("page"), {}).get("width", 1),
            "page_height": pages.get(region.get("page"), {}).get("height", 1),
        }
        for region in (protected_regions or [])
    ]
    # F26: a block is only "rendered" if its translation is readable in the
    # frame the engine used in the candidate PDF.
    from .render_verification import verify_placements

    verification_issues, verification = verify_placements(output, result.get("placements") or [])
    proof = measure_render_fidelity(
        render_source,
        str(output),
        source_ir=ir,
        translations=translations,
        protected_regions=regions,
        text_masks=result.get("text_masks"),
        work_dir=output.parent / f"{output.stem}-fidelity",
        # F27: targets verified in the engine frame are present even when the
        # PDF text extractor splits them over several blocks.
        verified_block_ids=set(verification.get("verified_block_ids") or ()),
    )
    issues = result.get("issues", []) + proof.pop("issues", [])
    result.update(proof)
    result["render_verification"] = verification
    result["issues"] = issues + verification_issues
    return result


def _reconstruct_docx(ir, translations, source, out, protected_regions, issues):
    from .docx_fidelity import edit_docx_package
    from .render_metrics import parse_docx_structure

    replacements = []
    for block in _all_blocks(ir):
        target = translations.get(block["id"], block.get("text", ""))
        if target == block.get("text", ""):
            continue
        if _protected_for(block, protected_regions):
            issues.append(
                _warn(
                    "PROTECTED_TRANSLATION_NOT_APPLIED",
                    "Zone originale protégée: traduction non appliquée.",
                    "error",
                    block_id=block["id"],
                )
            )
            continue
        replacements.append(
            {"id": block["id"], "source_ref": block.get("source_ref", {}), "target": target}
        )
    package = edit_docx_package(source, out, replacements)
    issues.extend(package["issues"])
    changed = package["changed_blocks"]
    candidate_structure = parse_docx_structure(out)
    candidate_work = out.parent / f"{out.stem}-render"
    candidate_pages, candidate_previews, render_warnings = _docx_rendered_pages(
        out, candidate_work, {}
    )
    issues.extend(render_warnings)
    return {
        "output_path": str(out),
        "issues": issues,
        "pages_preserved": len(candidate_pages) == len(ir.get("pages", []))
        if candidate_pages
        else None,
        "layout_score": None,
        "changed_blocks": changed,
        "pagination": "libreoffice-pdf" if candidate_pages else None,
        "candidate_pdf_path": str(candidate_work / "libreoffice" / f"{out.stem}.pdf")
        if candidate_pages
        else None,
        "candidate_preview_paths": [str(path) for path in candidate_previews],
        "renderer_version": "docx-package-fidelity-v3",
        "redistribution_trace": package["redistribution_trace"],
        "table_structure": [
            {
                "id": f"candidate-table-{t['index']}",
                "rows": t["rows"],
                "columns": t["columns"],
                "cell_count": len(t["cells"]),
                "provenance": "candidate_ooxml",
            }
            for t in candidate_structure.get("tables", [])
        ],
    }


def _reconstruct_office(ir, translations, source, out, protected_regions, issues):
    """RUN28 route A: rewrite text nodes only, then verify the package natively."""
    from .ooxml_office import VERSION as OFFICE_VERSION
    from .ooxml_office import edit_office_package, verify_office_candidate

    replacements = []
    for block in _all_blocks(ir):
        target = translations.get(block["id"], block.get("text", ""))
        if target == block.get("text", ""):
            continue
        if _protected_for(block, protected_regions):
            issues.append(
                _warn(
                    "PROTECTED_TRANSLATION_NOT_APPLIED",
                    "Zone originale protégée: traduction non appliquée.",
                    "error",
                    block_id=block["id"],
                )
            )
            continue
        if ir.get("format") == "pptx" and block.get("autofit") not in {"text", "shape"}:
            ratio = len(target) / max(1, len(block.get("text", "")))
            if ratio > 1.25:
                issues.append(
                    _warn(
                        "PPTX_TEXT_MAY_OVERFLOW",
                        "Traduction plus longue sans ajustement automatique : débordement possible (estimation, non mesuré).",
                        "warning",
                        block_id=block["id"],
                        length_ratio=round(ratio, 2),
                    )
                )
        replacements.append(
            {"id": block["id"], "source_ref": block.get("source_ref", {}), "target": target}
        )
    package = edit_office_package(source, out, replacements)
    issues.extend(package["issues"])
    expected = {
        r["id"]: r["target"] for r in replacements if r["id"] in set(package["changed_blocks"])
    }
    verification = verify_office_candidate(source, out, expected)
    not_applied = set(verification["not_applied"]) | {
        r["id"] for r in replacements if r["id"] not in set(package["changed_blocks"])
    }
    for block_id in sorted(not_applied):
        issues.append(
            _warn(
                "OFFICE_TARGET_NOT_APPLIED",
                "Traduction absente du fichier produit.",
                "critical",
                block_id=block_id,
            )
        )
    for block_id in verification["unexpected_changes"]:
        issues.append(
            _warn(
                "OFFICE_UNEXPECTED_TEXT_CHANGE",
                "Texte modifié hors traduction attendue.",
                "critical",
                block_id=block_id,
            )
        )
    if verification["parts_changed_beyond_text"] or not (
        verification["formulas_identical"]
        and verification["sheet_names_identical"]
        and verification["units_identical"]
    ):
        issues.append(
            _warn(
                "OFFICE_PACKAGE_INTEGRITY_FAILED",
                "Le paquet produit diffère de la source hors texte (formules, noms, structure).",
                "critical",
                parts=verification["parts_changed_beyond_text"][:20],
            )
        )
    candidate_previews = []
    pagination = None
    if ir.get("versions", {}).get("preview"):
        candidate_work = out.parent / f"{out.stem}-render"
        candidate_pages, candidate_previews, render_warnings = _docx_rendered_pages(
            out, candidate_work, {}
        )
        for warning in render_warnings:
            warning["code"] = warning["code"].replace("DOCX_", "OFFICE_")
            warning["severity"] = "warning"
        issues.extend(render_warnings)
        pagination = "libreoffice-pdf" if candidate_pages else None
    return {
        "output_path": str(out),
        "issues": issues,
        "pages_preserved": verification["units_identical"],
        "layout_score": None,
        "changed_blocks": sorted(set(package["changed_blocks"]) - not_applied),
        "pagination": pagination,
        "candidate_preview_paths": [str(path) for path in candidate_previews],
        "renderer_version": OFFICE_VERSION,
        "native_verification": verification,
        "rewritten_parts": package["rewritten_parts"],
        "placements": [],
    }


def reconstruct_document(
    ir: dict,
    translations: dict[str, str],
    output_path: str,
    protected_regions: list[dict] | None = None,
) -> dict:
    if not isinstance(translations, dict) or any(
        not isinstance(k, str) or not isinstance(v, str) for k, v in translations.items()
    ):
        raise ValueError("Translation candidates must be a mapping of text ids to text, never null")
    out = Path(output_path).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    source = Path(ir["source_path"])
    if source.resolve() == out.resolve() or (out.exists() and os.path.samefile(source, out)):
        raise ValueError("Original/source must not be overwritten by reconstruction")
    expected_sha = ir.get("versions", {}).get("source_sha256") or ir.get("versions", {}).get(
        "source_xml_sha256"
    )
    if expected_sha and hashlib.sha256(source.read_bytes()).hexdigest() != expected_sha:
        raise ValueError("Original immutable SHA256 mismatch")
    from .protected_visuals import normalize_regions

    protected_regions = normalize_regions(
        list(ir.get("protected_visual_regions") or []) + list(protected_regions or []),
        ir.get("pages", []),
        expected_sha or hashlib.sha256(source.read_bytes()).hexdigest(),
    )
    issues = []
    if ir.get("format") == "pdf":
        issues.append(
            _warn(
                "SIGNATURE_AUTHENTICITY_NOT_VERIFIED",
                "Le candidat dérivé conserve des éléments visuels mais ne porte pas de validation cryptographique de la signature source.",
            )
        )
    if ir.get("format") in OFFICE_FORMATS:
        return _reconstruct_office(ir, translations, source, out, protected_regions, issues)
    if ir.get("format") == "docx":
        result = _reconstruct_docx(ir, translations, source, out, protected_regions, issues)
        return _with_render_proof(ir, translations, out, result, protected_regions)
    from .pdf_fidelity import render_pdf

    render_source = source
    if ir.get("format") in {"tif", "tiff"}:
        from .raster_adapter import raster_pdf_copy

        render_source = raster_pdf_copy(source, out.parent / "source-raster-copy.pdf")
    result = render_pdf(ir, translations, render_source, out, protected_regions)
    result["issues"] = issues + result.get("issues", [])
    result["protected_regions"] = protected_regions
    result["renderer_version"] = "pdf-cell-fidelity-v3"
    return _with_render_proof(ir, translations, out, result, protected_regions)


def _quality_contract(ir, translations, render_result, reference, options=None):
    """Call the coordinator-owned quality contract without importing it at startup."""
    try:
        from importlib import import_module

        module = import_module("app.ai.quality.engine")
        evaluator = module.evaluate_quality
    except (ImportError, AttributeError):
        return None, _warn(
            "QUALITY_ENGINE_UNAVAILABLE",
            "Le moteur qualité contractuel est absent; aucun score qualité n'est fabriqué.",
            "error",
        )
    try:
        result = evaluator(ir, translations, render_result, reference, options or {})
    except Exception as exc:  # fail closed: a broken scorer is not a PASS
        return None, _warn("QUALITY_ENGINE_ERROR", str(exc), "error")
    if not isinstance(result, dict):
        return None, _warn(
            "QUALITY_ENGINE_INVALID_RESULT", "Résultat qualité non dictionnaire.", "error"
        )
    return result, None


def evaluate_document(
    ir: dict,
    translations: dict[str, str],
    render_result: dict | None = None,
    reference: dict | None = None,
    options: dict | None = None,
) -> dict:
    issues = []
    critical = []
    untranslated = []
    deltas = []
    blocks = _all_blocks(ir)
    known_ids = {block["id"] for block in blocks}
    for unknown_id in sorted(set(translations) - known_ids):
        issues.append(
            _warn(
                "UNKNOWN_TRANSLATION_BLOCK",
                "Identifiant de traduction absent de l'IR.",
                "error",
                block_id=unknown_id,
            )
        )
    protected_regions = (render_result or {}).get("protected_regions", [])
    for block in blocks:
        source = block.get("text", "")
        if block["id"] not in translations:
            issues.append(
                _warn(
                    "MISSING_TRANSLATION_BLOCK",
                    "Bloc source sans candidat de traduction; aucune coercition silencieuse.",
                    "error",
                    block_id=block["id"],
                )
            )
            continue
        target = translations[block["id"]]
        protected = _protected_for(block, protected_regions)
        # D27: a typography-only variant of the source is not a translation.
        if source.strip() and (target == source or typography_equivalent(target, source)):
            if protected:
                issues.append(
                    _warn(
                        "PROTECTED_TRANSLATION_NOT_APPLIED",
                        "Bloc protégé conservé; résultat explicitement inconnu.",
                        "warning",
                        block_id=block["id"],
                    )
                )
            elif _requires_translation(source) and (
                (block.get("language") and block.get("language") not in ("fra", "fr"))
                or (not block.get("language") and ir.get("language") not in ("fra", "fr"))
            ):
                untranslated.append(block["id"])
                issues.append(
                    _warn(
                        "UNTRANSLATED",
                        "Segment non français conservé sans traduction vérifiable.",
                        "error",
                        block_id=block["id"],
                    )
                )
        if target is not None:
            entity_issues = protected_entity_issues(source, target)
            for entity_issue in entity_issues:
                entity_issue["block_id"] = block["id"]
                if not ir.get("protected_entities") and entity_issue.get("entity_type") in {
                    "name",
                    "company",
                    "address",
                }:
                    entity_issue["severity"] = "warning"
                    entity_issue["code"] = "UNPROTECTED_ENTITY_REQUIRES_REVIEW"
            issues.extend(entity_issues)
            critical.extend(issue for issue in entity_issues if issue.get("severity") == "critical")
            deltas.append({"block_id": block["id"], "tokens": token_delta(source, target)})
    if render_result:
        issues.extend(render_result.get("issues", []))
    ref_valid = reference_is_valid(reference)
    geometry = (render_result or {}).get("layout_score")
    # Never expose raw pixel dissimilarity under a similarity/quality score.
    quality = {
        "ocr": None,
        "translation": None,
        "semantic": None,
        "reference_status": "validated" if ref_valid else "not_available",
        "geometry_only_layout": geometry,
        "visual_difference": (render_result or {}).get("visual_difference"),
        "visual_score_method": None,
        "numeric_entity_deltas": deltas,
    }
    contract_quality, contract_warning = _quality_contract(
        ir, translations, render_result, reference, options
    )
    if contract_warning:
        issues.append(contract_warning)
    result = {
        "critical_entity_errors": critical,
        "untranslated_blocks": untranslated,
        "layout_score": geometry,
        "visual_score": None,
        "ocr_score": None,
        "translation_score": None,
        "semantic_score": None,
        "issues": issues,
        "status": "REVIEW_REQUIRED",
        "quality": quality,
        "reference_status": "validated" if ref_valid else "not_evaluated",
        "metrics": reference_metrics(_source_text(ir), "\n".join(translations.values()), reference),
    }
    if contract_quality is not None:
        explicit_entities = bool(ir.get("protected_entities"))
        contract_issues = []
        for issue in contract_quality.get("issues", []):
            # Automatic name/company/address candidates are coverage warnings
            # until a caller supplies explicit protected annotations. Numeric
            # and identifier mismatches remain critical through the local
            # occurrence comparison above.
            if (
                not explicit_entities
                and isinstance(issue, dict)
                and issue.get("severity") == "critical"
                and issue.get("entity_type") in {"name", "company", "address"}
            ):
                issue = {
                    **issue,
                    "severity": "warning",
                    "code": "UNPROTECTED_ENTITY_REQUIRES_REVIEW",
                    "message": "Entité nominative détectée automatiquement sans annotation protégée explicite.",
                }
            contract_issues.append(issue)
        for key in (
            "ocr_score",
            "translation_score",
            "semantic_score",
            "layout_score",
            "visual_score",
            "status",
            "issues",
            "critical_entity_errors",
            "hallucination_errors",
            "reference_status",
            "provenance",
            "coverage",
        ):
            if key in contract_quality:
                result[key] = contract_quality[key]
        result["issues"] = issues + contract_issues
        result["untranslated_blocks"] = sorted(
            set(untranslated + (contract_quality.get("untranslated_blocks") or []))
        )
        result["critical_entity_errors"] = [
            issue
            for issue in result.get("critical_entity_errors", [])
            if issue.get("severity") == "critical"
            and (
                explicit_entities or issue.get("entity_type") not in {"name", "company", "address"}
            )
        ]
        result["hallucination_errors"] = [
            issue
            for issue in result.get("hallucination_errors", [])
            if explicit_entities or issue.get("entity_type") not in {"name", "company", "address"}
        ]
        result["quality_contract"] = contract_quality
    else:
        result["hallucination_errors"] = []
        result["provenance"] = {"quality_engine": "unavailable"}
        result["coverage"] = {"quality_engine": False}
    # Pipeline-facing reference lifecycle is distinct from provider availability.
    # Keep the provider's raw status in quality_contract/provenance; absence is
    # never a reference validation and never a PASS.
    if not reference:
        result["reference_status"] = "not_evaluated"
    if result.get("critical_entity_errors") or result.get("issues"):
        result["status"] = "REVIEW_REQUIRED"
    return result


def _normalise_typography_only(blocks: dict, translations: dict) -> list[dict]:
    """D27/F40 (run28): a target that differs from the source only by typography
    or by invisible bidi marks IS the source, whatever produced it (model,
    translation memory, another engine). Normalise it once, before
    reconstruction, so the renderer never re-typesets an OCR reading over the
    original pixels and every later check (render, verification, block status)
    sees it unchanged. Mutates ``translations``; returns one info issue per block.
    """
    issues = []
    for block_id, text in blocks.items():
        target = translations.get(block_id)
        if (
            isinstance(target, str)
            and target != text
            and str(text or "").strip()
            and typography_equivalent(target, text)
        ):
            translations[block_id] = text
            issues.append(
                _warn(
                    "TYPOGRAPHY_ONLY_TARGET_NORMALISED",
                    "Cible identique à la source à la typographie près : source conservée.",
                    "info",
                    block_id=block_id,
                )
            )
    return issues


async def translate_document(ir: dict, work_dir: str, options: dict | None = None) -> dict:
    options = options or {}
    work = Path(work_dir).resolve()
    work.mkdir(parents=True, exist_ok=True)
    (work / "previews").mkdir(parents=True, exist_ok=True)
    blocks = {block["id"]: block.get("text", "") for block in _all_blocks(ir)}
    issues = []
    result = {}
    try:
        from app.ai.translation.service import translate_blocks

        result = translate_blocks(ir, options)
        if asyncio.iscoroutine(result):
            result = await result
        result = result or {}
    except (ImportError, AttributeError):
        result = {}
        if options.get("engine", "offline") != "offline":
            issues.append(
                _warn(
                    "PRIVATE_TRANSLATION_UNAVAILABLE",
                    "Service de traduction privé indisponible; aucun appel public n'est effectué.",
                    "critical",
                )
            )
    except Exception as exc:
        result = {}
        issues.append(_warn("TRANSLATION_SERVICE_ERROR", str(exc), "critical"))
    translations = dict(result.get("translations") or {})
    if not translations:
        memory, glossary = {}, {}
        if options.get("engine", "offline") == "offline":
            try:
                from app.ai.translation.service import _approved_memory, _glossary

                memory = _approved_memory(options.get("memory"))
                glossary = _glossary(options.get("glossary"), ir.get("domain"))
            except (ImportError, AttributeError):
                pass
        for block_id, text in blocks.items():
            block = next(block for block in _all_blocks(ir) if block["id"] == block_id)
            translations[block_id] = memory.get(text, glossary.get(text, text))
            if (
                translations[block_id] == text
                and text.strip()
                and _requires_translation(text)
                and block.get("language") not in ("fra", "fr")
            ):
                issues.append(
                    _warn(
                        "UNTRANSLATED",
                        "Segment conservé hors ligne faute de mémoire/glossaire.",
                        "error",
                        block_id=block_id,
                    )
                )
    issues.extend(_normalise_typography_only(blocks, translations))
    untranslated = list(
        result.get("untranslated_blocks")
        or [
            block_id
            for block_id, text in blocks.items()
            if text.strip()
            and translations.get(block_id, text) == text
            and _requires_translation(text)
            and next(block for block in _all_blocks(ir) if block["id"] == block_id).get("language")
            not in ("fra", "fr")
        ]
    )
    issues.extend(result.get("issues") or [])
    source = Path(ir["source_path"])
    original_previews = sorted((work / "previews").glob("original-*.png"))
    if not original_previews and ir.get("format") == "pdf":
        original_previews = _preview_pdf(source, work, "original")
    if not original_previews and ir.get("format") in {"tif", "tiff"}:
        from .raster_adapter import raster_pdf_copy

        original_previews = _preview_pdf(
            raster_pdf_copy(source, work / "source-raster-copy.pdf"), work, "original"
        )
    native = ir.get("format") in OFFICE_FORMATS or ir.get("format") == "docx"
    candidate = work / (f"candidate.{ir.get('format')}" if native else "candidate.pdf")
    render = reconstruct_document(
        ir, translations, str(candidate), options.get("protected_regions")
    )
    if ir.get("format") in ("pdf", "jpg", "jpeg", "png", "tif", "tiff"):
        translated = _preview_pdf(candidate, work, "translated")
        for index, (before, after) in enumerate(
            zip(original_previews, translated, strict=False), 1
        ):
            try:
                raw_difference = _render_diff(
                    before, after, work / "previews" / f"diff-{index:03d}.png"
                )
                render.setdefault("raw_visual_differences", []).append(raw_difference)
            except Exception as exc:
                issues.append(_warn("DIFF_RENDER_ERROR", str(exc), "warning", page=index))
    else:
        translated = [Path(path) for path in render.get("candidate_preview_paths", [])]
        # DOCX previews are generated by the same isolated LibreOffice path as
        # the source/candidate pagination; never fabricate a page image here.
        for index, (before, after) in enumerate(
            zip(original_previews, translated, strict=False), 1
        ):
            try:
                _render_diff(before, after, work / "previews" / f"diff-{index:03d}.png")
            except Exception as exc:
                issues.append(_warn("DIFF_RENDER_ERROR", str(exc), "warning", page=index))
    render["source_preview_paths"] = [str(path) for path in original_previews]
    render["candidate_preview_paths"] = [str(path) for path in translated if Path(path).is_file()]
    render["diff_preview_paths"] = [
        str(path) for path in sorted((work / "previews").glob("diff-*.png"))
    ]
    render.setdefault("protected_regions", options.get("protected_regions") or [])
    render.setdefault("text_masks", [])
    if len(original_previews) != len(translated) and original_previews:
        issues.append(
            _warn(
                "PAGE_COUNT_CHANGED",
                "Le nombre de pages du candidat diffère du rendu source; revue requise.",
                "error",
                source_pages=len(original_previews),
                candidate_pages=len(translated),
            )
        )
    render["issues"] = render.get("issues", [])
    issues.extend(render["issues"])
    evaluation = evaluate_document(ir, translations, render, options.get("reference"), options)
    issues.extend(evaluation["issues"])
    from .fidelity_report import write_visual_artifacts

    fidelity = write_visual_artifacts(work, render, evaluation)
    versions = {
        "pipeline": "local-v2",
        "candidate": str(candidate),
        "renderer": render.get("renderer_version"),
    }
    service_versions = result.get("versions") or {}
    if isinstance(service_versions, dict):
        versions.update(service_versions)
    for key in ("model", "prompt", "glossary", "translation_memory"):
        if key in result:
            versions[key] = result[key]
    from .block_status import classify_blocks

    output = {
        "translations": translations,
        "issues": issues,
        "block_status": classify_blocks(ir, translations, issues),
        "quality": evaluation,
        "document_fidelity": fidelity,
        "versions": versions,
        "untranslated_blocks": untranslated,
        "translation_telemetry": result.get("translation_telemetry")
        or ir.get("_translation_telemetry")
        or {},
    }
    if render.get("native_verification"):
        # RUN28 route A: XLSX/PPTX native package verification (no visual proof).
        output["native_verification"] = render["native_verification"]
    _save_json(work / "translation.json", output)
    return output
