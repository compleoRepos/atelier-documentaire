from __future__ import annotations

import hashlib
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from app.documents.docx_analysis import collect_docx_content
from app.documents.docx_fidelity import edit_docx_package, redistribute_text
from app.documents.render_metrics import measure_render_fidelity

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML = "http://www.w3.org/XML/1998/namespace"
NS = {"w": W, "r": R}


def _zip_docx(path: Path, document: str, *, extra: dict[str, bytes] | None = None) -> None:
    members = {
        "[Content_Types].xml": b"synthetic content types",
        "word/document.xml": document.encode(),
        "word/media/keep.bin": b"media bytes must not change",
        "custom/unrelated.bin": bytes(range(64)),
    }
    members.update(extra or {})
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)


def _collect(path: Path) -> tuple[list[dict], list[dict], dict]:
    def make_block(text, block_id, order, source_ref):
        return {
            "id": block_id,
            "text": text,
            "page": 1,
            "language": "eng",
            "source_ref": source_ref,
        }

    return collect_docx_content(path, make_block)


def _text_nodes(xml: bytes) -> list[str]:
    root = etree.fromstring(xml)
    return [node.text or "" for node in root.xpath(".//w:t", namespaces=NS)]


def test_collect_docx_ir_retains_run_properties_and_numeric_source_ref(tmp_path: Path):
    document = f'''<w:document xmlns:w="{W}" xmlns:r="{R}" xmlns:xml="{XML}"><w:body>
      <w:p><w:pPr><w:jc w:val="both"/><w:ind w:left="720" w:firstLine="360"/><w:spacing w:before="120" w:after="80"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="7"/></w:numPr></w:pPr>
        <w:r><w:rPr><w:b/><w:i/><w:u w:val="single"/><w:color w:val="C00000"/><w:highlight w:val="yellow"/></w:rPr><w:t xml:space="preserve">Alpha </w:t></w:r>
        <w:r><w:rPr><w:sz w:val="28"/></w:rPr><w:t>Beta</w:t><w:tab/><w:br w:type="page"/></w:r>
      </w:p>
      <w:sectPr/>
    </w:body></w:document>'''
    source = tmp_path / "rich.docx"
    _zip_docx(source, document)

    blocks, tables, metadata = _collect(source)

    assert not tables
    assert len(blocks) == 1
    block = blocks[0]
    assert block["text"] == "Alpha Beta"
    assert block["source_ref"]["xml_part"] == "word/document.xml"
    assert isinstance(block["source_ref"]["paragraph_path"], list)
    assert all(isinstance(index, int) for index in block["source_ref"]["paragraph_path"])
    assert [run["text"] for run in block["runs"]] == ["Alpha ", "Beta"]
    assert "<w:b" in block["runs"][0]["properties_xml"]
    assert "<w:i" in block["runs"][0]["properties_xml"]
    assert "C00000" in block["runs"][0]["style_xml"]
    assert "yellow" in block["runs"][0]["style_xml"]
    assert "w:jc" in block["paragraph_properties_xml"]
    assert block["tabs"] and block["breaks"]
    assert block["properties"]["paragraph_style"] is None
    assert metadata["sections"]


def test_collect_docx_ir_tracks_merges_empty_cells_nested_targets_and_header_footer_parts(
    tmp_path: Path,
):
    document = f'''<w:document xmlns:w="{W}"><w:body>
      <w:p><w:r><w:t>Body target</w:t></w:r></w:p>
      <w:tbl><w:tblPr><w:tblW w:w="9000" w:type="dxa"/><w:tblBorders><w:top w:val="single"/></w:tblBorders><w:shd w:fill="D9EAF7"/></w:tblPr><w:tblGrid><w:gridCol w:w="3000"/><w:gridCol w:w="3000"/><w:gridCol w:w="3000"/></w:tblGrid>
        <w:tr><w:trPr><w:tblHeader/><w:trHeight w:val="420"/></w:trPr><w:tc><w:tcPr><w:gridSpan w:val="2"/><w:shd w:fill="FFFF00"/></w:tcPr><w:p><w:r><w:t>Horizontal</w:t></w:r></w:p></w:tc><w:tc><w:p/></w:tc></w:tr>
        <w:tr><w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr><w:p><w:r><w:t>Vertical</w:t></w:r></w:p><w:tbl><w:tblGrid><w:gridCol w:w="3000"/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>Nested target</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:tc><w:tc><w:p><w:r><w:t>Cell one</w:t></w:r></w:p><w:p><w:r><w:t>Cell two</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Right</w:t></w:r></w:p></w:tc></w:tr>
        <w:tr><w:tc><w:tcPr><w:vMerge w:val="continue"/></w:tcPr><w:p/></w:tc><w:tc><w:p><w:r><w:t>Bottom</w:t></w:r></w:p></w:tc><w:tc><w:p/></w:tc></w:tr>
      </w:tbl><w:sectPr/></w:body></w:document>'''
    headers = {
        "word/header1.xml": f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>First header</w:t></w:r></w:p></w:hdr>'.encode(),
        "word/header2.xml": f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>Even header</w:t></w:r></w:p></w:hdr>'.encode(),
        "word/footer1.xml": f'<w:ftr xmlns:w="{W}"><w:p><w:r><w:t>Shared footer</w:t></w:r></w:p></w:ftr>'.encode(),
        "word/footer2.xml": f'<w:ftr xmlns:w="{W}"><w:p><w:r><w:t>Even footer</w:t></w:r></w:p></w:ftr>'.encode(),
    }
    source = tmp_path / "hierarchy.docx"
    _zip_docx(source, document, extra=headers)

    blocks, tables, _metadata = _collect(source)

    assert {block["source_ref"]["xml_part"] for block in blocks} >= {
        "word/document.xml",
        "word/header1.xml",
        "word/header2.xml",
        "word/footer1.xml",
        "word/footer2.xml",
    }
    table = tables[0]
    assert table["row_records"] and table["cells"]
    assert table["grid_widths"] == [3000, 3000, 3000]
    assert any(cell["column_span"] == 2 for cell in table["cells"])
    assert any(cell["v_merge"] == "restart" for cell in table["cells"])
    continuation = next(cell for cell in table["cells"] if cell["v_merge"] == "continue")
    assert continuation["merge_parent_id"]
    assert any(cell["row_span"] == 2 for cell in table["cells"])
    assert any(not cell["paragraph_ids"] for cell in table["cells"])
    nested = next(cell for cell in table["cells"] if cell["table_ids"])
    assert nested["table_ids"]
    assert any(
        block["text"] == "Nested target"
        and block["source_ref"]["table_index"] > table["cells"][0]["source_ref"]["table_index"]
        for block in blocks
    )
    assert {block["text"] for block in blocks} >= {"Cell one", "Cell two"}


def test_multistyle_translation_preserves_ooxml_formatting_without_semantic_certification(
    tmp_path: Path,
):
    document = f'''<w:document xmlns:w="{W}" xmlns:xml="{XML}"><w:body><w:p>
      <w:pPr><w:jc w:val="center"/><w:ind w:left="720"/><w:spacing w:after="240"/><w:numPr><w:ilvl w:val="1"/><w:numId w:val="9"/></w:numPr></w:pPr>
      <w:r><w:rPr><w:b/><w:i/><w:u w:val="single"/><w:color w:val="1F4E79"/><w:highlight w:val="cyan"/></w:rPr><w:t xml:space="preserve">Source </w:t></w:r>
      <w:r><w:rPr><w:sz w:val="26"/><w:color w:val="C00000"/></w:rPr><w:t>words</w:t></w:r>
    </w:p></w:body></w:document>'''
    source = tmp_path / "styled.docx"
    output = tmp_path / "styled-translated.docx"
    _zip_docx(source, document)
    before_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    root_before = etree.fromstring(document.encode())
    before_styles = [etree.tostring(node) for node in root_before.xpath(".//w:rPr", namespaces=NS)]
    before_ppr = etree.tostring(root_before.xpath(".//w:pPr", namespaces=NS)[0])
    blocks, _tables, _metadata = _collect(source)
    ref = blocks[0]["source_ref"]

    result = edit_docx_package(
        source, output, [{"id": blocks[0]["id"], "source_ref": ref, "target": "Mots traduits"}]
    )

    assert result["status"] == "REVIEW_REQUIRED"
    assert result["changed_blocks"] == 1
    assert result["replacements"][0]["status"] == "REVIEW_REQUIRED"
    assert result["replacements"][0]["trace"]["semantic_alignment"] is False
    assert output != source
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before_sha
    with zipfile.ZipFile(output) as archive:
        translated = etree.fromstring(archive.read("word/document.xml"))
        assert "Mots traduits" == "".join(_text_nodes(archive.read("word/document.xml")))
        assert [
            etree.tostring(node) for node in translated.xpath(".//w:rPr", namespaces=NS)
        ] == before_styles
        assert etree.tostring(translated.xpath(".//w:pPr", namespaces=NS)[0]) == before_ppr
        assert archive.read("word/media/keep.bin") == b"media bytes must not change"
        assert archive.read("custom/unrelated.bin") == bytes(range(64))
        assert b'xml:space="preserve"' in archive.read("word/document.xml")


def test_package_edit_uses_numeric_nested_and_header_refs_and_rejects_output_alias(tmp_path: Path):
    document = f'''<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>Body target</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:tbl><w:tr><w:tc><w:p><w:r><w:t>Nested target</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:tc></w:tr></w:tbl></w:body></w:document>'''
    source = tmp_path / "refs.docx"
    output = tmp_path / "refs-out.docx"
    _zip_docx(
        source,
        document,
        extra={
            "word/header1.xml": f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>Header target</w:t></w:r></w:p></w:hdr>'.encode()
        },
    )
    blocks, _tables, _metadata = _collect(source)
    by_text = {block["text"]: block for block in blocks}
    replacements = [
        {"id": "body", "source_ref": by_text["Body target"]["source_ref"], "target": "Corps"},
        {
            "id": "nested",
            "source_ref": by_text["Nested target"]["source_ref"],
            "target": "Imbrique",
        },
        {"id": "header", "source_ref": by_text["Header target"]["source_ref"], "target": "En-tete"},
    ]

    result = edit_docx_package(source, output, replacements)

    assert result["status"] == "APPLIED"
    assert result["changed_blocks"] == 3
    assert set(result["modified_parts"]) == {"word/document.xml", "word/header1.xml"}
    with zipfile.ZipFile(output) as archive:
        assert _text_nodes(archive.read("word/document.xml")) == ["Corps", "Imbrique"]
        assert _text_nodes(archive.read("word/header1.xml")) == ["En-tete"]
        assert archive.read("word/media/keep.bin") == b"media bytes must not change"
    with pytest.raises(ValueError):
        edit_docx_package(source, source, replacements)


def test_hyperlink_bookmark_field_tab_and_break_translation_is_review_only_and_preserved():
    paragraph = etree.fromstring(
        f'''<w:p xmlns:w="{W}" xmlns:r="{R}"><w:bookmarkStart w:id="1" w:name="anchor"/><w:hyperlink r:id="rId7"><w:r><w:t>Linked source</w:t><w:tab/><w:br/></w:r></w:hyperlink><w:bookmarkEnd w:id="1"/><w:r><w:fldChar w:fldCharType="begin"/><w:instrText xml:space="preserve"> PAGE </w:instrText><w:fldChar w:fldCharType="separate"/><w:t>4</w:t><w:fldChar w:fldCharType="end"/></w:r></w:p>'''
    )
    before = etree.tostring(paragraph)

    result = redistribute_text(paragraph, "Traduction")

    assert result["status"] == "REVIEW_REQUIRED"
    assert result["applied"] is False
    assert etree.tostring(paragraph) == before
    codes = {issue["code"] for issue in result["issues"]}
    assert {
        "DOCX_HYPERLINK_BOUNDARY_AMBIGUOUS",
        "DOCX_BOOKMARK_BOUNDARY_AMBIGUOUS",
        "DOCX_FIELD_RESULT_PROTECTED",
        "DOCX_LAYOUT_MARKER_AMBIGUOUS",
    } <= codes
    assert paragraph.xpath(".//w:t", namespaces=NS)[0].text == "Linked source"
    assert paragraph.xpath(".//w:t", namespaces=NS)[1].text == "4"


def test_real_libreoffice_pagination_change_is_reported_for_review(tmp_path: Path):
    if shutil.which("soffice") is None:
        pytest.skip("LibreOffice is required for the real pagination integration check")
    source = tmp_path / "one-page.docx"
    candidate = tmp_path / "many-pages.docx"
    source_doc = Document()
    source_doc.add_paragraph("Short source paragraph.")
    source_doc.save(source)
    candidate_doc = Document()
    for index in range(90):
        candidate_doc.add_paragraph(
            f"Translated paragraph {index}: " + "longer translated content " * 18
        )
    candidate_doc.save(candidate)

    result = measure_render_fidelity(source, candidate, work_dir=tmp_path / "render")

    assert result["provenance"]["source_pdf"] == "libreoffice-pdf"
    assert result["source_page_count"] != result["candidate_page_count"]
    assert any(issue["code"] == "RENDER_PAGE_COUNT_CHANGED" for issue in result["issues"])
