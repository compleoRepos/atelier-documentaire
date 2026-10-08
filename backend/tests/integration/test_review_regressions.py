from __future__ import annotations

import io
import json

import fitz
from docx import Document as WordDocument

from app.api.routes import _entity_ok
from app.core.settings import get_settings
from tests.conftest import login


def synthetic_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((40, 60), "Attestation de travail 1234", fontsize=12)
    page.insert_text((40, 120), "Nous attestons que le salarié est en activité.", fontsize=12)
    content = document.tobytes()
    document.close()
    return content


def prepared_document(client, headers):
    response = client.post(
        "/api/documents",
        headers=headers,
        files={"file": ("synthetic.pdf", synthetic_pdf(), "application/pdf")},
    )
    assert response.status_code == 200, response.text
    identity = response.json()["id"]
    response = client.post(f"/api/documents/{identity}/translate", headers=headers)
    assert response.status_code == 200, response.text
    details = client.get(f"/api/documents/{identity}", headers=headers).json()
    assert details["status"] == "REVIEW_REQUIRED"
    return identity, details


def test_entity_occurrences_are_compared_only_where_they_exist():
    assert _entity_ok("Attestation", "Justificatif", ["Nom fictif"])
    assert _entity_ok("Nom fictif: Alice", "Nom fictif : Alice", ["Alice"])
    assert not _entity_ok("Alice Alice", "Alice", ["Alice"])
    assert not _entity_ok("Attestation", "Attestation Alice", ["Alice"])


def test_approval_cannot_silently_ignore_unsaved_corrections(client):
    headers = login(client, "reviewer-a", "Reviewer-A-pass-12")
    identity, details = prepared_document(client, headers)
    block = details["analysis"]["pages"][0]["blocks"][0]["id"]
    response = client.post(
        f"/api/documents/{identity}/review",
        headers=headers,
        json={
            "action": "approve",
            "attest_review": True,
            "expected_version": details["review_version"],
            "corrections": {block: "Justificatif de travail 1234"},
        },
    )
    assert response.status_code == 409
    assert (
        client.get(f"/api/documents/{identity}", headers=headers).json()["status"]
        == "REVIEW_REQUIRED"
    )


def test_stale_review_version_is_rejected_and_detail_keeps_protections(client):
    headers = login(client, "reviewer-a", "Reviewer-A-pass-12")
    identity, details = prepared_document(client, headers)
    response = client.post(
        f"/api/documents/{identity}/review",
        headers=headers,
        json={
            "action": "correct",
            "expected_version": details["review_version"],
            "corrections": {},
            "protected_entities": ["le salarié"],
        },
    )
    assert response.status_code == 200, response.text
    current = client.get(f"/api/documents/{identity}", headers=headers).json()
    assert current["review_version"] > details["review_version"]
    assert current["protected_entities"] == ["le salarié"]
    assert "untranslated_blocks" in current["translation"]
    response = client.post(
        f"/api/documents/{identity}/review",
        headers=headers,
        json={
            "action": "approve",
            "attest_review": True,
            "expected_version": details["review_version"],
        },
    )
    assert response.status_code == 409


def test_global_and_other_tenant_qa_reports_are_not_exposed(client):
    headers_a = login(client, "admin-a", "Admin-A-password-12")
    headers_b = login(client, "user-b", "User-B-password-12")
    tenant_a = client.get("/api/auth/me", headers=headers_a).json()["tenant_id"]
    report_dir = get_settings().report_dir
    report_dir.mkdir(parents=True, exist_ok=True)
    global_report = {"documents": 99, "tests_passed": 99, "tests_failed": 0, "metrics": {}}
    (report_dir / "summary.json").write_text(json.dumps(global_report), encoding="utf-8")
    assert client.get("/api/testing", headers=headers_b).json()["documents"] is None
    scoped = report_dir / tenant_a
    scoped.mkdir()
    (scoped / "summary.json").write_text(
        json.dumps({**global_report, "tenant_id": tenant_a}), encoding="utf-8"
    )
    assert client.get("/api/testing", headers=headers_a).json()["documents"] == 99
    assert client.get("/api/testing", headers=headers_b).json()["documents"] is None


def test_docx_correction_does_not_use_pdf_renderer(client):
    headers = login(client, "reviewer-a", "Reviewer-A-pass-12")
    word = WordDocument()
    word.add_paragraph("Employment Certificate")
    content = io.BytesIO()
    word.save(content)
    response = client.post(
        "/api/documents",
        headers=headers,
        files={
            "file": (
                "synthetic.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200, response.text
    identity = response.json()["id"]
    assert client.post(f"/api/documents/{identity}/translate", headers=headers).status_code == 200
    details = client.get(f"/api/documents/{identity}", headers=headers).json()
    block = details["analysis"]["pages"][0]["blocks"][0]["id"]
    response = client.post(
        f"/api/documents/{identity}/review",
        headers=headers,
        json={
            "action": "correct",
            "expected_version": details["review_version"],
            "corrections": {block: "Attestation de travail"},
        },
    )
    assert response.status_code == 200, response.text
    candidate = client.get(f"/api/documents/{identity}/artifact/candidate", headers=headers)
    assert candidate.status_code == 200 and candidate.content.startswith(b"PK")
    corrected = WordDocument(io.BytesIO(candidate.content))
    assert corrected.paragraphs[0].text == "Attestation de travail"


def test_previous_preview_remains_identical_after_new_review_version(client):
    headers = login(client, "reviewer-a", "Reviewer-A-pass-12")
    identity, details = prepared_document(client, headers)
    old_version = details["review_version"]
    old_preview = client.get(
        f"/api/documents/{identity}/preview/translated/1?version={old_version}", headers=headers
    )
    assert old_preview.status_code == 200
    block = details["analysis"]["pages"][0]["blocks"][0]["id"]
    response = client.post(
        f"/api/documents/{identity}/review",
        headers=headers,
        json={
            "action": "correct",
            "expected_version": old_version,
            "corrections": {block: "Justificatif de travail 1234"},
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["review_version"] > old_version
    unchanged = client.get(
        f"/api/documents/{identity}/preview/translated/1?version={old_version}", headers=headers
    )
    assert unchanged.content == old_preview.content
    assert (
        client.get(
            f"/api/documents/{identity}/preview/translated/1?version=999", headers=headers
        ).status_code
        == 404
    )


def test_human_golden_translation_enters_validated_memory_and_is_evaluation_reference(client):
    headers = login(client, "reviewer-a", "Reviewer-A-pass-12")
    word = WordDocument()
    word.add_paragraph("Employment Certificate")
    content = io.BytesIO()
    word.save(content)
    upload = client.post(
        "/api/documents",
        headers=headers,
        files={
            "file": (
                "synthetic-golden.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 200
    identity = upload.json()["id"]
    assert client.post(f"/api/documents/{identity}/translate", headers=headers).status_code == 200
    detail = client.get(f"/api/documents/{identity}", headers=headers).json()
    block = detail["analysis"]["pages"][0]["blocks"][0]["id"]
    correction = client.post(
        f"/api/documents/{identity}/review",
        headers=headers,
        json={
            "action": "correct",
            "expected_version": detail["review_version"],
            "corrections": {block: "Attestation de travail"},
        },
    )
    assert correction.status_code == 200, correction.text
    approval = client.post(
        f"/api/documents/{identity}/review",
        headers=headers,
        json={
            "action": "approve",
            "expected_version": correction.json()["review_version"],
            "attest_review": True,
            "promote_golden": True,
        },
    )
    assert approval.status_code == 200, approval.text
    memory = client.get("/api/memory", headers=headers).json()
    assert any(
        entry["source"] == "Employment Certificate"
        and entry["target"] == "Attestation de travail"
        and entry["validated"]
        for entry in memory
    )
    assert client.post(f"/api/documents/{identity}/translate", headers=headers).status_code == 200
    repeated = client.get(f"/api/documents/{identity}", headers=headers).json()
    assert repeated["quality"]["reference_status"] == "validated"
    assert repeated["quality"]["metrics"]["translation_cer"] == 0.0
    assert repeated["quality"]["metrics"]["cer"] is None
