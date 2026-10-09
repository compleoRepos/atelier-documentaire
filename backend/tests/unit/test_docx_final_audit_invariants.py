from __future__ import annotations

import hashlib
import posixpath
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from lxml import etree
from PIL import Image

from app.documents.docx_analysis import collect_docx_content
from app.documents.docx_fidelity import edit_docx_package, paragraph_metadata, redistribute_text
from app.documents.pipeline import _convert_docx_with_libreoffice

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "r": R, "pr": PR}


def _collect(path: Path):
    return collect_docx_content(
        path,
        lambda text, bid, order, ref: {
            "id": bid,
            "text": text,
            "page": 1,
            "source_ref": ref,
        },
    )


def _add_hyperlink(paragraph, document, text: str):
    rid = document.part.relate_to("https://example.test/audit", RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _build_fixture(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "En-tête source"
    section.footer.paragraphs[0].text = "Pied source"

    paragraph = doc.add_paragraph()
    paragraph.add_run("Normal ")
    bold = paragraph.add_run("bold ")
    bold.bold = True
    italic = paragraph.add_run("italic")
    italic.italic = True
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Après saut de page")

    listed = doc.add_paragraph("Élément liste", style="List Bullet")
    ppr = listed._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.extend([ilvl, num_id])
    ppr.append(num_pr)

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cellule A"
    table.cell(0, 1).text = "Cellule B"
    table.cell(1, 0).text = "Cellule C"
    table.cell(1, 1).text = "Cellule D"

    link_paragraph = doc.add_paragraph()
    _add_hyperlink(link_paragraph, doc, "Lien source")

    image_path = path.parent / "fixture.png"
    Image.new("RGB", (12, 12), (20, 80, 140)).save(image_path)
    doc.add_picture(str(image_path))
    doc.save(path)

    # Add explicit merged-cell properties while retaining all cell text.
    with zipfile.ZipFile(path) as archive:
        members = {info.filename: archive.read(info) for info in archive.infolist()}
    root = etree.fromstring(members["word/document.xml"])
    first_table = root.xpath(".//w:tbl", namespaces=NS)[0]
    rows = first_table.xpath("./w:tr", namespaces=NS)
    first_cell = rows[0].xpath("./w:tc", namespaces=NS)[0]
    first_props = first_cell.find(f"{{{W}}}tcPr")
    grid_span = etree.Element(f"{{{W}}}gridSpan")
    grid_span.set(f"{{{W}}}val", "2")
    first_props.append(grid_span)
    restart_cell = rows[1].xpath("./w:tc", namespaces=NS)[0]
    restart_props = restart_cell.find(f"{{{W}}}tcPr")
    restart = etree.Element(f"{{{W}}}vMerge")
    restart.set(f"{{{W}}}val", "restart")
    restart_props.append(restart)
    members["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)


def _texts(root):
    return [node.text or "" for node in root.xpath(".//w:t", namespaces=NS)]


def test_first_run_rule_preserves_each_style_when_word_count_and_order_change():
    paragraph = etree.fromstring(
        f"<w:p {nsdecls('w')}>"
        '<w:r><w:rPr><w:rFonts w:ascii="Arial"/><w:sz w:val="22"/></w:rPr><w:t xml:space="preserve">Normal </w:t></w:r>'
        '<w:r><w:rPr><w:b/><w:color w:val="C00000"/></w:rPr><w:t xml:space="preserve">bold </w:t></w:r>'
        '<w:r><w:rPr><w:i/><w:u w:val="single"/></w:rPr><w:t>italic</w:t></w:r>'
        "</w:p>"
    )
    before_styles = [run.find(f"{{{W}}}rPr") for run in paragraph]
    result = redistribute_text(paragraph, "Texte réordonné et beaucoup plus long")
    runs = paragraph.xpath("./w:r", namespaces=NS)
    assert result["applied"] is True
    assert "".join(_texts(paragraph)) == "Texte réordonné et beaucoup plus long"
    assert all(run.xpath("string(w:t)", namespaces=NS) for run in runs)
    assert [etree.tostring(run.find(f"{{{W}}}rPr")) for run in runs] == [
        etree.tostring(style) for style in before_styles
    ]
    assert len(runs[0].xpath("string(w:t)", namespaces=NS)) < len(
        "Texte réordonné et beaucoup plus long"
    )
    assert any(issue["code"] == "DOCX_MULTI_STYLE_BOUNDARY_AMBIGUOUS" for issue in result["issues"])


def test_synthetic_package_keeps_ooxml_structures_relationships_and_editability(tmp_path):
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _build_fixture(source)
    blocks, tables, analysis = _collect(source)
    assert tables and tables[0]["cells"]
    assert analysis["sections"]
    by_text = {block["text"]: block for block in blocks}
    replacements = [
        {
            "id": by_text["Normal bold italic"]["id"],
            "source_ref": by_text["Normal bold italic"]["source_ref"],
            "target": "Normal traduit gras italique",
        },
        {
            "id": by_text["Lien source"]["id"],
            "source_ref": by_text["Lien source"]["source_ref"],
            "target": "Lien traduit",
        },
        {
            "id": by_text["En-tête source"]["id"],
            "source_ref": by_text["En-tête source"]["source_ref"],
            "target": "En-tête traduit",
        },
        {
            "id": by_text["Pied source"]["id"],
            "source_ref": by_text["Pied source"]["source_ref"],
            "target": "Pied traduit",
        },
    ]
    result = edit_docx_package(source, candidate, replacements)
    # The styled paragraph also contains an explicit page break and is
    # correctly refused as ambiguous; the hyperlink, header and footer apply.
    assert result["changed_count"] == 3
    assert all(item["status"] in {"APPLIED", "REVIEW_REQUIRED"} for item in result["replacements"])

    with zipfile.ZipFile(source) as before, zipfile.ZipFile(candidate) as after:
        before_xml = etree.fromstring(before.read("word/document.xml"))
        after_xml = etree.fromstring(after.read("word/document.xml"))
        assert after_xml.xpath("string(.//w:hyperlink/w:r/w:t)", namespaces=NS) == "Lien traduit"
        assert after_xml.xpath(".//w:numPr", namespaces=NS)
        assert etree.tostring(after_xml.xpath(".//w:numPr", namespaces=NS)[0]) == etree.tostring(
            before_xml.xpath(".//w:numPr", namespaces=NS)[0]
        )
        assert etree.tostring(after_xml.xpath(".//w:gridSpan", namespaces=NS)[0]) == etree.tostring(
            before_xml.xpath(".//w:gridSpan", namespaces=NS)[0]
        )
        assert etree.tostring(after_xml.xpath(".//w:vMerge", namespaces=NS)[0]) == etree.tostring(
            before_xml.xpath(".//w:vMerge", namespaces=NS)[0]
        )
        assert after_xml.xpath(".//w:sectPr", namespaces=NS)
        assert after_xml.xpath('.//w:br[@w:type="page"]', namespaces=NS)
        assert after.read("word/_rels/document.xml.rels") == before.read(
            "word/_rels/document.xml.rels"
        )
        assert b"En-t\xc3\xaate traduit" in after.read("word/header1.xml")
        assert b"Pied traduit" in after.read("word/footer1.xml")
        assert [name for name in after.namelist() if name.startswith("word/media/")] == [
            name for name in before.namelist() if name.startswith("word/media/")
        ]
        for name in before.namelist():
            if name.startswith("word/media/"):
                assert (
                    hashlib.sha256(after.read(name)).digest()
                    == hashlib.sha256(before.read(name)).digest()
                )
        for rel_name in (
            "word/_rels/document.xml.rels",
            "word/_rels/header1.xml.rels",
            "word/_rels/footer1.xml.rels",
        ):
            if rel_name not in after.namelist():
                continue
            rels = etree.fromstring(after.read(rel_name))
            base = posixpath.dirname(rel_name.removeprefix("word/_rels/").removesuffix(".rels"))
            owner = posixpath.basename(rel_name.removeprefix("word/_rels/").removesuffix(".rels"))
            for rel in rels:
                target = rel.get("Target")
                if rel.get("TargetMode") == "External":
                    continue
                resolved = posixpath.normpath(posixpath.join("word", base, target))
                assert resolved in after.namelist(), (owner, target, resolved)

    opened = Document(candidate)
    assert "Normal bold italic" in " ".join(p.text for p in opened.paragraphs)
    assert "Lien traduit" in " ".join(p.text for p in opened.paragraphs)
    assert zipfile.is_zipfile(candidate)


def test_source_and_candidate_pagination_are_measured_with_libreoffice(tmp_path):
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _build_fixture(source)
    blocks, _, _ = _collect(source)
    body = next(block for block in blocks if block["text"] == "Après saut de page")
    result = edit_docx_package(
        source,
        candidate,
        [
            {
                "id": body["id"],
                "source_ref": body["source_ref"],
                "target": "Après saut de page traduit",
            }
        ],
    )
    assert result["changed_count"] == 1
    source_pdf = _convert_docx_with_libreoffice(source, tmp_path / "source-pdf")
    candidate_pdf = _convert_docx_with_libreoffice(candidate, tmp_path / "candidate-pdf")
    with fitz.open(source_pdf) as source_doc, fitz.open(candidate_pdf) as candidate_doc:
        assert source_doc.page_count == 2
        assert candidate_doc.page_count == source_doc.page_count


def test_paragraph_metadata_exposes_rpr_and_real_text_not_images():
    paragraph = parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:rPr><w:b/><w:i/><w:u w:val="single"/><w:color w:val="00AA00"/><w:sz w:val="28"/></w:rPr><w:t>Style</w:t></w:r></w:p>'
    )
    metadata = paragraph_metadata(paragraph)
    assert metadata["runs"][0]["style_xml"]
    assert all(
        token in metadata["runs"][0]["style_xml"] for token in ("b", "i", "u", "color", "sz")
    )
    assert metadata["text"] == "Style"
