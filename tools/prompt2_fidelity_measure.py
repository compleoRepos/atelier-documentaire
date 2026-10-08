"""Private measured artifacts. Synthetic targets are renderer fixtures, not Golden/Qwen."""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import importlib.util
import json
import subprocess
import time
from pathlib import Path

import fitz
from app.documents import pipeline
from app.documents.fidelity_report import write_visual_artifacts
from app.documents.ooxml_integrity import compare_ooxml_integrity
from docx import Document
from docx.shared import Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
BASE = "6fb913a3a402e45c5813bf343fb85f26ef0e0e20"


def legacy_pipeline(work):
    # Load the actual base implementation, not a handcrafted imitation.
    text = subprocess.check_output(
        ["git", "show", f"{BASE}:backend/app/documents/pipeline.py"], cwd=ROOT
    )
    source = work / "baseline_pipeline.py"
    source.write_bytes(text)
    spec = importlib.util.spec_from_file_location("app.documents.baseline_prompt2", source)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def synthetic(work):
    work.mkdir(parents=True, exist_ok=True)
    old = legacy_pipeline(work)
    source = work / "styled.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Employee ").bold = True
    paragraph.add_run("salary ").italic = True
    run = paragraph.add_run("information")
    run.font.size = Pt(12)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Header source"
    table.cell(1, 0).text = "Cell source"
    table.cell(1, 1).text = "1234"
    doc.sections[0].header.paragraphs[0].text = "Header text"
    doc.sections[0].footer.paragraphs[0].text = "Footer text"
    doc.save(source)
    before_ir = old.analyze_document(str(source), str(work / "docx-before-analysis"))
    after_ir = pipeline.analyze_document(str(source), str(work / "docx-after-analysis"))
    memory = {
        "Employee salary information": "Informations salariales employé",
        "Header source": "En-tête source",
        "Cell source": "Texte cellule",
        "Header text": "Texte en-tête",
        "Footer text": "Texte pied",
    }
    docx_results = {}
    for label, engine, ir in [("before", old, before_ir), ("after", pipeline, after_ir)]:
        translations = {
            b["id"]: memory.get(b["text"], b["text"]) for p in ir["pages"] for b in p["blocks"]
        }
        out = work / f"docx-{label}.docx"
        render = engine.reconstruct_document(ir, translations, str(out))
        evaluation = engine.evaluate_document(ir, translations, render)
        artifacts = write_visual_artifacts(work / f"docx-{label}", render, evaluation)
        candidate = Document(out)
        docx_results[label] = {
            "changed_blocks": render["changed_blocks"],
            "nonempty_runs": sum(bool(r.text) for r in candidate.paragraphs[0].runs),
            "run_count": len(candidate.paragraphs[0].runs),
            "run_texts": [r.text for r in candidate.paragraphs[0].runs],
            "package_invariants": compare_ooxml_integrity(source, out)["preserved"],
            "pages": render["candidate_page_count"],
            "layout_score": render["layout_score"],
            "semantic_style_score": render["structure_metrics"]["styles"]["score"],
            "issue_codes": sorted({i["code"] for i in render["issues"]}),
            "artifacts": artifacts["pages"],
        }
    source_pdf = work / "table.pdf"
    image = work / "original-graphic.png"
    Image.new("RGB", (32, 32), "blue").save(image)
    with fitz.open() as doc:
        page = doc.new_page(width=320, height=220)
        for row in range(2):
            for col in range(2):
                rect = fitz.Rect(25 + col * 130, 35 + row * 55, 155 + col * 130, 90 + row * 55)
                page.draw_rect(rect, color=(0, 0, 0.8), fill=(0.94, 0.94, 0.65), width=1.2)
                page.insert_text(
                    (rect.x0 + 7, rect.y0 + 25), f"Cell source {row}{col}", fontsize=11
                )
        page.insert_image(fitz.Rect(250, 170, 282, 202), filename=str(image))
        doc.save(source_pdf)
    pdf_results = {}
    for label, engine in [("before", old), ("after", pipeline)]:
        ir = engine.analyze_document(str(source_pdf), str(work / f"pdf-{label}-analysis"))
        translations = {
            b["id"]: "Texte cellule " + b["text"][-2:] for p in ir["pages"] for b in p["blocks"]
        }
        out = work / f"pdf-{label}.pdf"
        render = engine.reconstruct_document(
            ir,
            translations,
            str(out),
            [{"page": 1, "bbox": [250, 170, 282, 202], "subtype": "image_candidate"}],
        )
        evaluation = engine.evaluate_document(ir, translations, render)
        artifacts = write_visual_artifacts(work / f"pdf-{label}", render, evaluation)
        from app.documents.pdf_fidelity import compare_pdf_tables

        measured = compare_pdf_tables(source_pdf, out, source_ir=ir, translations=translations)
        pdf_results[label] = {
            "changed_blocks": render["changed_blocks"],
            "overflow_count": render.get("overflow_count"),
            "tables": measured["structure_metrics"],
            "layout_score": render["layout_score"],
            "raster_similarity": render["raster_similarity"],
            "protected_pixels": render["protected_region_results"],
            "cell_render_units": len(render.get("table_render_trace", [])),
            "artifacts": artifacts["pages"],
        }
    summary = {
        "schema": "synthetic-renderer-measurement/v1",
        "golden_reference": False,
        "live_qwen": False,
        "docx": docx_results,
        "pdf": pdf_results,
    }
    (work / "measurements.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("Synthetic before/after measured", flush=True)
    return summary


def corpus(work):
    catalog = json.loads((ROOT / "dataset/metadata/catalog.json").read_text())
    results = []
    for entry in catalog["documents"]:
        source = ROOT / "dataset/input" / entry["filename"]
        sha = hashlib.sha256(source.read_bytes()).hexdigest()
        if sha != entry["sha256"]:
            raise ValueError("Corpus original SHA mismatch")
        docwork = work / entry["document_id"]
        started = time.monotonic()
        print(f"Analyzing {entry['document_id']} ({entry['pages']} pages)", flush=True)
        ir = pipeline.analyze_document(str(source), str(docwork))
        result = asyncio.run(pipeline.translate_document(ir, str(docwork), {"engine": "offline"}))
        unchanged = sha == hashlib.sha256(source.read_bytes()).hexdigest()
        record = {
            "document_id": entry["document_id"],
            "pages": len(ir["pages"]),
            "original_immutable": unchanged,
            "status": result["quality"]["status"],
            "reference_status": result["quality"]["reference_status"],
            "translation_score": result["quality"]["translation_score"],
            "semantic_score": result["quality"]["semantic_score"],
            "live_qwen": False,
            "human_validated": False,
            "seconds": round(time.monotonic() - started, 2),
            "table_candidates": len(ir.get("tables", [])),
            "protected_visual_candidates": len(ir.get("protected_visual_regions", [])),
            "untranslated_count": len(result["untranslated_blocks"]),
            "artifact_page_count": len(result["document_fidelity"]["pages"]),
        }
        results.append(record)
        (work / "corpus-results.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
        print(
            f"Measured {entry['document_id']}: {record['status']}, immutable={unchanged}",
            flush=True,
        )
    return results


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=ROOT / "test-report/prompt-2/measured")
    parser.add_argument("--corpus", action="store_true")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    if args.corpus:
        corpus(args.out)
    else:
        synthetic(args.out)
